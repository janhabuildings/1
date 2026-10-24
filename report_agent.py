"""Template-preserving inspection report generation; never sends mail."""
import argparse
from copy import deepcopy
from datetime import date
from email.message import EmailMessage
from email.policy import SMTP
from html import escape
import json
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from zipfile import ZipFile
from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parent
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
W = '{' + NS['w'] + '}'


def require(value, name):
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f'Missing or invalid {name}')
    if any(ord(c) < 32 and c not in '\n\t' for c in value):
        raise ValueError(f'Invalid control character in {name}')
    return value.strip()


def validate(data, *, draft=False):
    if data.get('equipment') not in ('SF', 'SH', 'OT'):
        raise ValueError('equipment must be SF, SH or OT')
    if data.get('status') not in ('OIs', 'No OIs'):
        raise ValueError('status must be OIs or No OIs; do not infer from missing findings')
    for name in ('client', 'project', 'requested_by', 'permit', 'address', 'locations', 'inspection_code'):
        require(data.get(name), name)
    for name in ('request_date', 'inspection_date'):
        date.fromisoformat(require(data.get(name), name))
    for name in ('drawings_on_site', 'drawings_approved'):
        if data.get(name) not in (('Yes', 'No', 'NA', '[Pending confirmation]') if draft else ('Yes', 'No', 'NA')):
            raise ValueError(f'{name} must be Yes, No, or NA')
    if data.get('hoist') or data.get('chute'):
        raise ValueError('Hoist/chute inspections require a separate supported template')
    attendees = data.get('attendees')
    if not isinstance(attendees, list) or not attendees:
        raise ValueError('Provide actual attendees, including the inspecting PE')
    for attendee in attendees:
        for name in ('name', 'company', 'title'):
            require(attendee.get(name), 'attendee.' + name)
    if not any(a['title'].strip().upper() == 'PE' for a in attendees):
        raise ValueError('An inspecting PE attendee must be provided')
    items = data.get('open_items')
    if not isinstance(items, list):
        raise ValueError('open_items must be an explicit list')
    if bool(items) != (data['status'] == 'OIs'):
        raise ValueError('Open items conflict with report status')
    for item in items:
        for name in ('location', 'action'):
            require(item.get(name), 'open_item.' + name)
    if data['equipment']=='OT': require(product(data),'Other item name')
    fields = () if data['equipment']=='OT' else ('length_ft', 'height_ft', 'sections', 'width_ft') if data['equipment'] == 'SH' else (
        'length_ft', 'height_ft', 'tie_horizontal', 'tie_vertical', 'planked_levels',
        'outrigger_levels', 'tie_type', 'dunnage_type')
    for name in fields:
        require(data.get('site', {}).get(name), 'site.' + name)
    if draft:
        return  # Report-only review draft: no mail headers or recipient claims.
    for name in ('sender', 'to', 'cc'):
        values = data.get(name)
        values = [values] if name == 'sender' else values
        if not isinstance(values, list) or not values:
            raise ValueError(f'{name} is required; use verified client recipients and inspections mailbox')
        for address in values:
            if not isinstance(address, str) or not re.fullmatch(r'[^\s<>@,;]+@[^\s<>@,;]+\.[^\s<>@,;]+', address):
                raise ValueError(f'Invalid email in {name}')
    if data['sender'].casefold() not in [x.casefold() for x in data['cc']]:
        raise ValueError('CC must include the DSENY inspections sender mailbox')
    if data.get('recipients_confirmed') is not True:
        raise ValueError('Confirm recipients belong to the client; never copy calendar attendees blindly')


def replace_cell(cell, value):
    """Keep all paragraph/run formatting, replacing only text and inline content."""
    p = cell.paragraphs[0]
    exemplar = next((r for r in p.runs if r.text.strip()), p.runs[0] if p.runs else None)
    props = deepcopy(exemplar._r.rPr) if exemplar is not None and exemplar._r.rPr is not None else None
    for paragraph in list(cell._tc.findall(W + 'p'))[1:]:
        cell._tc.remove(paragraph)
    for child in list(p._p):
        if child.tag != W + 'pPr':
            p._p.remove(child)
    run = p.add_run(str(value))
    if props is not None:
        run._r.insert(0, props)


def set_rows(table, start, entries):
    exemplar = deepcopy(table.rows[start]._tr)
    for row in list(table.rows)[start:]:
        table._tbl.remove(row._tr)
    for entry in entries:
        table._tbl.append(deepcopy(exemplar))
        for cell, value in zip(table.rows[-1].cells, entry):
            replace_cell(cell, value)


# Supported stable placeholders; unknown fields are never silently dropped.
TEMPLATE_FIELDS = tuple(('client project subject_project requested_by permit address locations inspection_code '
    'request_date inspection_date drawings_on_site drawings_approved product status attendees open_items '
    'source_references photo_references').split()) + tuple('site.'+k for k in (
    'length_ft','height_ft','sections','width_ft','tie_horizontal','tie_vertical',
    'planked_levels','outrigger_levels','tie_type','dunnage_type','report_type','other_name','details'))
TOKEN = re.compile(r'\{\{\s*([A-Za-z_][A-Za-z0-9_.]*)\s*\}\}')


def template_values(data):
    result={key:str(data.get(key,'')) for key in TEMPLATE_FIELDS if not key.startswith('site.')}
    for key in ('request_date','inspection_date'):
        result[key]=date.fromisoformat(data[key]).strftime('%m.%d.%Y')
    result['product']=product(data)
    result['attendees']='\n'.join(' | '.join(a[k] for k in ('name','company','title')) for a in data['attendees'])
    result['open_items']='\n'.join(f'{i}. {x["location"]}: {x["action"]}' for i,x in enumerate(data['open_items'],1)) or 'No open items.'
    for key in ('source_references','photo_references'): result[key]='\n'.join(data.get(key,[]))
    for key in TEMPLATE_FIELDS:
        if key.startswith('site.'): result[key]=str(data.get('site',{}).get(key[5:],'[Pending confirmation]'))
    return result


def fill_placeholders(data, output, source, mapping):
    values=template_values(data)
    parser=etree.XMLParser(resolve_entities=False,load_dtd=False,no_network=True)
    with ZipFile(source) as original, ZipFile(output,'w') as result:
        for item in original.infolist():
            content=original.read(item.filename)
            if item.filename.startswith('word/') and item.filename.endswith('.xml'):
                root=etree.fromstring(content,parser)
                changed=False
                for paragraph in root.iter(W+'p'):
                    # Tokens may be split across Word runs. Replace from the right so offsets remain stable.
                    nodes=list(paragraph.iter(W+'t'))
                    text=''.join(n.text or '' for n in nodes)
                    offsets=[];position=0
                    for node in nodes:
                        offsets.append((position,position+len(node.text or ''),node));position=offsets[-1][1]
                    matches=list(TOKEN.finditer(text))
                    if matches: changed=True
                    for match in reversed(matches):
                        key=mapping.get(match[1],match[1])
                        if key not in values: raise ValueError('Unmapped template field: '+match[1])
                        first=next(i for i,(a,b,_) in enumerate(offsets) if a<=match.start()<b)
                        last=next(i for i,(a,b,_) in enumerate(offsets) if a<match.end()<=b)
                        a,_,node=offsets[first];prefix=(node.text or '')[:match.start()-a]
                        end_a,_,end_node=offsets[last];suffix=(end_node.text or '')[match.end()-end_a:]
                        node.text=prefix+values[key]+(suffix if first==last else '')
                        node.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
                        # Convert inserted line breaks to Word line-break nodes without changing run properties.
                        if '\n' in (node.text or ''):
                            pieces=node.text.split('\n');node.text=pieces[0]
                            parent=node.getparent();index=parent.index(node)+1
                            for piece in pieces[1:]:
                                parent.insert(index,etree.Element(W+'br'));index+=1
                                next_node=etree.Element(W+'t');next_node.text=piece
                                next_node.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
                                parent.insert(index,next_node);index+=1
                        if first!=last:
                            for _,_,middle in offsets[first+1:last]: middle.text=''
                            end_node.text=suffix
                if changed: content=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
            result.writestr(item,content)
    return source


def fill_report(data, output, *, draft=False, template=None, template_mode="legacy", mapping=None):
    validate(data, draft=draft)
    variant = 'Ois' if data['status'] == 'OIs' else 'N-Ois'
    source = ROOT / 'templates' / ('New Template Full.docx' if data['equipment']=='OT' else f'New Template {data["equipment"]}-{variant}.docx')
    if template is not None: source=Path(template)
    if template_mode == "placeholders":
        return fill_placeholders(data,output,source,mapping or {})
    if template_mode != "legacy": raise ValueError("Unsupported template mode")
    doc = Document(source)
    if data['equipment']=='OT': prepare_other_layout(doc,data)
    table = doc.tables[1]
    values = {(1,1): data['client'], (1,4): data['project'],
              (2,1): date.fromisoformat(data['request_date']).strftime('%m.%d.%Y'),
              (2,3): date.fromisoformat(data['inspection_date']).strftime('%m.%d.%Y'),
              (2,5): data['requested_by'], (3,1): product(data),
              (4,1): data['permit'], (5,1): data['address'],
              (6,3): data['drawings_on_site'], (7,3): data['drawings_approved']}
    for (row, col), value in values.items():
        replace_cell(table.cell(row, col), value)
    set_rows(doc.tables[2], 1, [('', a['name'], a['company'], a['title']) for a in data['attendees']])
    site = doc.tables[3]
    # Actual physical cells avoid aliasing merged grid cells.
    rows = [list(row.cells) for row in site.rows]
    loc_index = next(i for i,r in enumerate(rows) if r[0].text.strip() == 'LOCATION/S:')
    replace_cell(site.cell(loc_index, 1), data['locations'])
    measures = data['site']
    if data['equipment']=='OT':
        detail=deepcopy(site.rows[loc_index]._tr);site._tbl.append(detail)
        replace_cell(site.rows[-1].cells[0],'Description / scope:')
        replace_cell(site.rows[-1].cells[1],data.get('site',{}).get('details') or 'Not provided')
    elif data['equipment'] == 'SH':
        for col, key in ((1,'length_ft'),(3,'sections'),(5,'height_ft'),(7,'width_ft')):
            replace_cell(site.cell(loc_index + 1, col), measures[key])
    else:
        physical = []
        for cell in site.rows[loc_index+1].cells:
            if not physical or cell._tc is not physical[-1]._tc:
                physical.append(cell)
        for col,key in ((1,'length_ft'),(3,'height_ft'),(5,'tie_horizontal'),(7,'tie_vertical')):
            replace_cell(physical[col], measures[key])
        replace_cell(site.cell(loc_index+2, 2), measures['planked_levels'])
        replace_cell(site.cell(loc_index+2, 7 if variant=='Ois' else 6), measures['outrigger_levels'])
        replace_cell(site.cell(loc_index+3, 2), measures['tie_type'])
        replace_cell(site.cell(loc_index+4, 2), measures['dunnage_type'])
        if variant == 'Ois':
            # These example rows are not present in the No OIs master.
            for row in list(site.rows)[loc_index+5:]:
                site._tbl.remove(row._tr)
    if data['status'] == 'OIs':
        t = doc.tables[4]
        start = next(i for i,row in enumerate(t.rows) if row.cells[0].text.strip() == 'Number') + 1
        set_rows(t, start, [(str(i), item['location'], item['action']) for i,item in enumerate(data['open_items'],1)])
        for row in doc.tables[5].rows:
            replace_cell(row.cells[1], '')
    elif data['equipment']=='OT':
        t=doc.tables[4]
        replace_cell(t.rows[0].cells[0],'Section 4 Inspection findings — No open items')
        set_rows(t,2,[('',data['locations'],'No open items were reported in the inspected locations.')])
        doc.tables[5]._element.getparent().remove(doc.tables[5]._element)
    # Copy all original package parts verbatim except the populated main document.
    xml = etree.tostring(doc._element, xml_declaration=True, encoding='UTF-8', standalone=True)
    with ZipFile(source) as original, ZipFile(output, 'w') as result:
        for item in original.infolist():
            result.writestr(item, xml if item.filename == 'word/document.xml' else original.read(item.filename))
    return source


def prepare_other_layout(doc,data):
    # Use only the Full master's general sections, never its instruction text or example deficiencies.
    site=doc.tables[3];rows=list(site.rows)
    start=next(i for i,r in enumerate(rows) if ' '.join(r.cells[0].text.split())=='Section 3 Location')
    for row in rows:
        if row not in rows[start:start+2]: site._tbl.remove(row._tr)
    findings=doc.tables[4];rows=list(findings.rows)
    start=next(i for i,r in enumerate(rows) if ' '.join(r.cells[0].text.split())=='Section 4 Open Items')
    for row in rows[:start]: findings._tbl.remove(row._tr)
    body=doc._element.body;children=list(body)
    first=children.index(doc.tables[0]._element);second=children.index(doc.tables[1]._element)
    for index,child in enumerate(children):
        if child.tag!=W+'p': continue
        text=''.join(child.itertext()).strip()
        keep=first<index<second or (data['status']=='OIs' and any(marker in text for marker in (
            'Section 5, Client certification','Upon completion of corrected open items','I have personally and visually inspected')))
        if not keep: body.remove(child)


def product(data):
    if data['equipment']=='OT':
        site=data.get('site',{});kind=site.get('report_type','Unknown / custom')
        return require(site.get('other_name') if kind in ('','Unknown / custom','Other','Unknown') else kind,'Other item name')
    return 'Supported Scaffold' if data['equipment'] == 'SF' else 'Sidewalk Shed'


def subject(data):
    return f'{data["inspection_date"]} {data["project"]} {product(data)}-{data["inspection_code"]}-Inspection-{data["status"]}'


def convert(docx, directory, kind):
    office = shutil.which('libreoffice') or shutil.which('soffice')
    if not office and (ROOT/'.tools/bin/libreoffice').is_file():
        office = str(ROOT/'.tools/bin/libreoffice')
    if not office:
        raise RuntimeError('LibreOffice Writer is required for PDF/HTML conversion; install it or convert with Microsoft Word')
    with tempfile.TemporaryDirectory(prefix='inspection-office-') as profile:
        result = subprocess.run([office, '-env:UserInstallation=' + Path(profile).as_uri(), '--headless',
                                 '--convert-to', kind, '--outdir', str(directory), str(docx)],
                                capture_output=True, text=True, timeout=120)
    suffix = 'html' if kind.startswith('html') else 'pdf'
    output = directory / (docx.stem + '.' + suffix)
    if result.returncode or not output.is_file() or not output.stat().st_size:
        raise RuntimeError(f'Office conversion failed: {result.stderr or result.stdout}')
    if suffix == 'pdf' and not output.read_bytes().startswith(b'%PDF-'):
        raise RuntimeError('Converter did not produce a PDF')
    return output


def build(data, directory):
    validate(data)
    directory.mkdir(parents=True, exist_ok=True)
    title = subject(data)
    safe = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', title)
    target = directory / safe
    if target.exists():
        raise ValueError('Output already exists; use a new directory to preserve earlier drafts')
    target.mkdir()
    docx = target / (safe + '.docx')
    fill_report(data, docx)
    manifest = {'subject': title, 'status': 'report_generated', 'docx': docx.name,
                'outlook_draft_created': False, 'sent': False, 'delivery_ready': False,
                'rendering_note': 'Exact template fonts and Word/Outlook appearance require verification; conversion may substitute fonts.',
                'review_required': ['Check fonts, table layout and page breaks in Word/Outlook',
                                    'Verify source facts, inspection findings and recipients']}
    manifest_file = target / 'manifest.json'
    try:
        msg = EmailMessage(policy=SMTP)
        msg['Subject'], msg['From'], msg['To'], msg['Cc'] = title, data['sender'], ', '.join(data['to']), ', '.join(data['cc'])
        msg['X-Unsent'] = '1'
        if data['status'] == 'OIs':
            html_path = convert(docx, target, 'html:HTML (StarWriter)')
            html = html_path.read_text(encoding='utf-8')
            # Avoid an email that references local images inaccessible to recipients.
            if re.search(r'<img\b', html, re.I):
                raise RuntimeError('HTML contains images; inline/CID embedding is needed before email packaging')
            msg.set_content('Please view this inspection report in an HTML-capable email client.')
            msg.add_alternative(html, subtype='html')
        else:
            pdf = convert(docx, target, 'pdf')
            # Conversion never constitutes professional signing/sealing.
            text = ('Dear Client,\n\nEQ inspection report provided here. There are no open items. '
                    'Certification for use is included in the conclusion in the attached PDF.\n'
                    'See attached; please print & keep at the jobsite. It is the responsibility of the '
                    'contractor to ensure these reports are forwarded to all appropriate parties.\n\n')
            text += 'Regards,\n' + next(a['name'] for a in data['attendees'] if a['title'].strip().upper() == 'PE')
            msg.set_content(text)
            msg.add_alternative('<html><body>' + ''.join('<p>'+escape(p).replace('\n','<br>')+'</p>' for p in text.split('\n\n')) + '</body></html>', subtype='html')
            msg.add_attachment(pdf.read_bytes(), maintype='application', subtype='pdf', filename=pdf.name)
            manifest['pdf'] = pdf.name
            manifest['review_required'].append('PE must sign/seal the PDF; update cover wording only after signature verification')
        eml = target / (safe + '.eml')
        eml.write_bytes(bytes(msg))
        manifest.update(status='local_draft_generated', eml=eml.name)
    except (RuntimeError, subprocess.TimeoutExpired) as error:
        manifest.update(status='conversion_blocked', error=str(error))
    manifest_file.write_text(json.dumps(manifest, indent=2))
    return manifest_file, manifest


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('--output', type=Path, default=ROOT/'output')
    args = parser.parse_args()
    try:
        path, result = build(json.loads(args.input.read_text()), args.output.resolve())
        print(path)
        if result['status'] == 'conversion_blocked':
            parser.exit(2, result['error'] + '\nDOCX retained; no complete email draft produced.\n')
    except (ValueError, KeyError) as error:
        parser.exit(2, str(error) + '\n')
