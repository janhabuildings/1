"""Derive report-only variants from the reviewed October 25, 2026 full master."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.oxml import OxmlElement
from report_agent import W,physical_cells

ROOT=Path(__file__).resolve().parents[1]
REVISION='DSENY-2026-10-25'


def build():
    source=ROOT/'templates/New Template Full.docx'
    for equipment,site_index,items_index,cert_index in [('SH',3,5,6),('SF',7,8,9),('OT',10,12,13)]:
        for status in ('Ois','N-Ois'):
            doc=Document(source);tables=list(doc.tables);paragraphs=list(doc.paragraphs)
            body=doc._element.body
            section=deepcopy(body.sectPr)
            for node in list(body):body.remove(node)
            def append(node):body.append(deepcopy(node))
            if status=='Ois':
                for index in (44,45,47,48):append(paragraphs[index]._p)
            append(tables[0]._tbl);append(paragraphs[56]._p)
            append(tables[1]._tbl);append(tables[2]._tbl)
            site=deepcopy(tables[site_index]._tbl)
            for row in list(site.findall(W+'tr'))[{'SH':3,'SF':12,'OT':2}[equipment]:]:site.remove(row)
            body.append(site)
            if status=='Ois':
                # The generic Other report retains its narrower scope, rather than
                # adopting scaffold/shed structural certification as a finding.
                if equipment!='OT':
                    for index in ({'SH':(90,91),'SF':(104,105)}[equipment]):append(paragraphs[index]._p)
                items=deepcopy(tables[items_index]._tbl)
                for row in list(items.findall(W+'tr')):
                    text=''.join(row.itertext()).strip()
                    if 'Section 4' in text:break
                    items.remove(row)
                body.append(items)
                for index in {'SH':(95,96,98),'SF':(108,109,111),'OT':(125,126,128)}[equipment]:append(paragraphs[index]._p)
                cert=deepcopy(tables[cert_index]._tbl)
                for row in list(cert.findall(W+'tr'))[3:]:cert.remove(row)
                body.append(cert)
            else:
                append(tables[4]._tbl)
                parts=paragraphs[69:84] if equipment=='SH' else physical_cells(tables[7 if equipment=='SF' else 11].rows[13 if equipment=='SF' else 0])[0].paragraphs[4:19]
                for paragraph in parts:
                    if paragraph.text.strip():append(paragraph._p)
            body.append(section)
            doc.core_properties.identifier=REVISION
            rendered=BytesIO();doc.save(rendered);rendered.seek(0)
            destination=ROOT/'templates'/f'New Template {equipment}-{status}.docx'
            with ZipFile(source) as original,ZipFile(rendered) as updated,ZipFile(destination,'w') as output:
                for item in original.infolist():
                    output.writestr(item,updated.read(item.filename) if item.filename in ('word/document.xml','docProps/core.xml') else original.read(item.filename))
            print(destination.name)

if __name__=='__main__':build()
