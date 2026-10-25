"""Private reusable visual stamp assets; placement never creates a cryptographic signature."""
from io import BytesIO
import hashlib
from pathlib import Path
import re
from PIL import Image
from backend.core import uid
from report_agent import is_pe,example_stamp_rel_ids,strip_example_stamp_nodes


def initialize(store):
    (store.root/'signing').mkdir(exist_ok=True)
    with store.connect() as db:
        db.execute('CREATE TABLE IF NOT EXISTS signing_assets(id TEXT PRIMARY KEY,owner TEXT NOT NULL,sha256 TEXT NOT NULL,path TEXT NOT NULL)')


def save(store,content,owner):
    if not owner.strip() or len(owner)>200: raise ValueError('Enter the engineer name associated with this signature/stamp')
    if len(content)>4*1024*1024: raise ValueError('Stamp image limit is 4 MB')
    try:
        with Image.open(BytesIO(content)) as im:
            if im.format not in ('PNG','JPEG') or im.width*im.height>16000000: raise ValueError('Use a PNG/JPEG image up to 16 megapixels')
            im.verify()
    except Exception as error: raise ValueError('Invalid stamp image; use PNG or JPEG') from error
    digest=hashlib.sha256(content).hexdigest()
    with store.connect() as db:
        row=db.execute('SELECT id FROM signing_assets WHERE sha256=? AND owner=?',(digest,owner.strip())).fetchone()
        if row:return {'id':row['id'],'owner':owner.strip(),'sha256':digest}
    identity=uid();path=store.root/'signing'/(identity+'.image');path.write_bytes(content)
    with store.connect() as db:db.execute('INSERT INTO signing_assets VALUES (?,?,?,?)',(identity,owner.strip(),digest,str(path)))
    return {'id':identity,'owner':owner.strip(),'sha256':digest}


def get(store,identity):
    with store.connect() as db: row=db.execute('SELECT * FROM signing_assets WHERE id=?',(identity,)).fetchone()
    if not row: raise KeyError('Signing image not found')
    record=dict(row)
    if not Path(record['path']).is_file(): raise KeyError('Signing image file not available')
    content=Path(record['path']).read_bytes()
    if hashlib.sha256(content).hexdigest()!=record['sha256']: raise ValueError('Signing image checksum changed')
    return record


def validate_owner(asset,reports):
    normalize=lambda text:re.sub(r'\W+','',text.casefold())
    for report in reports:
        if report['status']!='No OIs': continue
        if not any(is_pe(a.get('title','')) and normalize(a.get('name',''))==normalize(asset['owner']) for a in report.get('attendees',[])):
            raise ValueError('The saved stamp engineer must match an actual PE attendee on each No OIs report')


def apply(pdf,asset):
    from pypdf import PdfReader,PdfWriter
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib.utils import ImageReader
    reader=PdfReader(pdf);writer=PdfWriter()
    for page in reader.pages:writer.add_page(page)
    first=writer.pages[0]
    if len(writer.pages)<2:writer.add_blank_page(width=float(first.mediabox.width),height=float(first.mediabox.height))
    page=writer.pages[1];width=float(page.mediabox.width);height=float(page.mediabox.height)
    buffer=BytesIO();canvas=Canvas(buffer,pagesize=(width,height))
    canvas.drawImage(ImageReader(asset['path']),width-235,height-214,width=180,height=144,preserveAspectRatio=True,anchor='c',mask='auto')
    canvas.setFont('Helvetica',8);canvas.drawRightString(width-55,height-224,'PE signature/seal: '+asset['owner'])
    canvas.save();buffer.seek(0);page.merge_page(PdfReader(buffer).pages[0],over=False)
    output=pdf.with_suffix('.stamping.tmp')
    with output.open('wb') as stream:writer.write(stream)
    output.replace(pdf)
    return {'asset_id':asset['id'],'sha256':asset['sha256'],'owner':asset['owner'],'page':2,'layer':'behind_text','cryptographic_signature':False}


def apply_docx(docx,asset):
    """Embed the original stamp as a floating Word Behind Text image in Section 3."""
    from copy import deepcopy
    from zipfile import ZipFile
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.shared import Pt
    doc=Document(docx)
    strip_example_stamp_nodes(doc._element,example_stamp_rel_ids(doc))
    paragraph=next((table.cell(0,0).paragraphs[0] for table in doc.tables
                    if re.search(r'Section\s*3\b',table.cell(0,0).text,re.I)),None)
    if paragraph is None:
        paragraph=next((p for p in doc.paragraphs if re.search(r'Section\s*3\b',p.text,re.I)),None)
    if paragraph is None:raise ValueError('A Section 3 heading is required to anchor the stamp in the DOCX')
    with Image.open(asset['path']) as image:iw,ih=image.size
    scale=min(180/iw,144/ih);width=iw*scale;height=ih*scale
    inline=paragraph.add_run().add_picture(asset['path'],width=Pt(width),height=Pt(height))._inline
    anchor=OxmlElement('wp:anchor')
    for key,value in {'distT':'0','distB':'0','distL':'0','distR':'0','simplePos':'0','relativeHeight':'0','behindDoc':'1','locked':'0','layoutInCell':'0','allowOverlap':'1'}.items():anchor.set(key,value)
    simple=OxmlElement('wp:simplePos');simple.set('x','0');simple.set('y','0');anchor.append(simple)
    for axis,offset in [('H',doc.sections[0].page_width.pt-235+(180-width)/2),('V',70+(144-height)/2)]:
        position=OxmlElement('wp:position'+axis);position.set('relativeFrom','page')
        amount=OxmlElement('wp:posOffset');amount.text=str(int(Pt(offset)));position.append(amount);anchor.append(position)
    anchor.append(deepcopy(inline.extent))
    effect=OxmlElement('wp:effectExtent')
    for key in ('l','t','r','b'):effect.set(key,'0')
    anchor.append(effect);anchor.append(OxmlElement('wp:wrapNone'))
    for child in inline:
        if child.tag.rsplit('}',1)[-1] in ('docPr','cNvGraphicFramePr','graphic'):
            copied=deepcopy(child)
            if child.tag.rsplit('}',1)[-1]=='docPr':copied.set('descr','PE signature/seal: '+asset['owner'])
            anchor.append(copied)
    inline.getparent().replace(inline,anchor)
    rendered=BytesIO();doc.save(rendered);rendered.seek(0)
    # Keep all unrelated original template package parts byte-for-byte unchanged.
    output=docx.with_suffix('.stamping.tmp')
    with ZipFile(docx) as original,ZipFile(rendered) as updated,ZipFile(output,'w') as result:
        changed={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
        old_names=set(original.namelist())
        for item in original.infolist():result.writestr(item,updated.read(item.filename) if item.filename in changed else original.read(item.filename))
        for item in updated.infolist():
            if item.filename not in old_names:result.writestr(item,updated.read(item.filename))
    output.replace(docx)
    return {'asset_id':asset['id'],'sha256':asset['sha256'],'owner':asset['owner'],'anchor':'Section 3','layer':'behind_text','cryptographic_signature':False}
