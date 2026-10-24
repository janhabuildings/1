"""Versioned DOCX templates. Files and metadata live under the configured data directory."""
from datetime import datetime,timezone
from io import BytesIO
import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile,BadZipFile
from lxml import etree
from docx import Document
from report_agent import TEMPLATE_FIELDS,TOKEN,NS
from backend.core import uid

SLOTS=('SF-OIs','SF-No OIs','SH-OIs','SH-No OIs','OT-OIs','OT-No OIs')
ROOT=Path(__file__).resolve().parents[1]


def inspect(content,mode,mapping,slot,*,builtin=False):
    if mode not in ('legacy','placeholders'): raise ValueError('Choose legacy or placeholders mode')
    if len(content)>10*1024*1024: raise ValueError('DOCX upload limit is 10 MB')
    if not isinstance(mapping,dict) or len(mapping)>100 or any(
        not isinstance(k,str) or not isinstance(v,str) or v not in TEMPLATE_FIELDS or not TOKEN.fullmatch('{{'+k+'}}')
        for k,v in mapping.items()): raise ValueError('Mappings must connect placeholder names to supported fields')
    try:
        with ZipFile(BytesIO(content)) as z:
            names=z.namelist()
            if len(names)!=len(set(names)) or len(names)>2000 or sum(i.file_size for i in z.infolist())>40*1024*1024:
                raise ValueError('Invalid or oversized DOCX package')
            if '[Content_Types].xml' not in names or 'word/document.xml' not in names: raise ValueError('Upload a Word DOCX file')
            if any('..' in Path(n).parts or n.startswith('/') or any(p in n.lower() for p in ('vbaproject','activex','embeddings/','_xmlsignatures')) for n in names):
                raise ValueError('Macros, embedded objects, signatures and unsafe package paths are unsupported')
            fields=set();invalid_marker=False;parser=etree.XMLParser(resolve_entities=False,load_dtd=False,no_network=True)
            for name in names:
                if not name.endswith(('.xml','.rels')): continue
                raw=z.read(name)
                root=etree.fromstring(raw,parser)
                if root.getroottree().docinfo.doctype: raise ValueError('Document type declarations are unsupported')
                if b'macroEnabled' in raw: raise ValueError('Macro-enabled templates are unsupported')
                if name.endswith('.rels'):
                    for rel in root:
                        if not builtin and rel.get('TargetMode')=='External' and not rel.get('Type','').endswith('/hyperlink'):
                            raise ValueError('Linked external images or objects are unsupported')
                if name.startswith('word/'):
                    instructions=root.xpath('//w:instrText/text() | //w:fldSimple/@w:instr',namespaces=NS)
                    if root.findall('.//w:altChunk',NS) or any(re.search(r'\b(DDEAUTO|DDE|INCLUDETEXT|INCLUDEPICTURE|DATABASE|LINK)\b',i,re.I) for i in instructions):
                        raise ValueError('External-content fields and embedded HTML are unsupported')
                    for paragraph in root.findall('.//w:p',NS):
                        text=''.join(paragraph.xpath('.//w:t/text()',namespaces=NS))
                        fields.update(TOKEN.findall(text))
                        stripped=TOKEN.sub("",text)
                        if "{{" in stripped or "}}" in stripped: invalid_marker=True
    except (BadZipFile,etree.XMLSyntaxError,RuntimeError,KeyError) as error:
        raise ValueError('Invalid DOCX package') from error
    errors=[]
    if invalid_marker: errors.append("Malformed placeholder: keep each {{field_name}} within one paragraph.")
    if mode=='legacy':
        try:
            doc=Document(BytesIO(content))
            # Exercise the positional mapping before accepting layout-compatible revisions.
            if len(doc.tables)!=(6 if slot.startswith('OT') or slot.endswith('-OIs') else 4): raise ValueError()
            if slot.startswith('OT'):
                site_headings=[' '.join(r.cells[0].text.split()) for r in doc.tables[3].rows]
                item_headings=[' '.join(r.cells[0].text.split()) for r in doc.tables[4].rows]
                if 'Section 3 Location' not in site_headings or 'Section 4 Open Items' not in item_headings: raise ValueError()
            site_rows=len(doc.tables[3].rows)
            if slot.startswith('SF') and site_rows<6: raise ValueError()
            if slot.startswith('SH') and site_rows!=3: raise ValueError()
            for row,col in ((1,1),(1,4),(2,5),(7,3)): doc.tables[1].cell(row,col)
            if not any(row.cells[0].text.strip()=='LOCATION/S:' for row in doc.tables[3].rows): raise ValueError()
        except (ValueError,IndexError,KeyError):
            raise ValueError('Legacy mode requires the original DSENY table layout. Use placeholders mode for a redesigned layout.')
        if fields: errors.append('Legacy mode cannot contain placeholders; select placeholders mode.')
    else:
        if not fields: errors.append('No {{field}} placeholders found.')
        unknown=sorted(k for k in fields if mapping.get(k,k) not in TEMPLATE_FIELDS)
        if unknown: errors.append('Map unsupported fields: '+', '.join(unknown))
        represented={mapping.get(k,k) for k in fields}
        required={'client','project','product','inspection_date','status','attendees','locations','open_items'}
        missing=sorted(required-represented)
        if missing: errors.append('Required report fields missing: '+', '.join(missing))
    return {'fields':sorted(fields),'errors':errors,'valid':not errors}


class TemplateManager:
    def __init__(self,store):
        self.store=store;self.directory=store.root/'templates';self.directory.mkdir(exist_ok=True)
        with store.connect() as db:
            db.executescript('''
            CREATE TABLE IF NOT EXISTS template_versions(
                id TEXT PRIMARY KEY,slot TEXT NOT NULL,label TEXT NOT NULL,mode TEXT NOT NULL,
                mapping TEXT NOT NULL,sha256 TEXT NOT NULL,created_at TEXT NOT NULL,builtin INTEGER NOT NULL DEFAULT 0,
                previewed INTEGER NOT NULL DEFAULT 0);
            CREATE TABLE IF NOT EXISTS template_archives(id TEXT PRIMARY KEY REFERENCES template_versions(id));
            CREATE TABLE IF NOT EXISTS template_artifacts(id TEXT PRIMARY KEY,template_id TEXT NOT NULL REFERENCES template_versions(id),
                filename TEXT NOT NULL,mime TEXT NOT NULL,path TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS template_active(slot TEXT PRIMARY KEY,version_id TEXT NOT NULL REFERENCES template_versions(id));
            ''')
        for slot in SLOTS:
            variant='N-Ois' if slot.endswith('-No OIs') else 'Ois'
            source=ROOT/'templates'/('New Template Full.docx' if slot.startswith('OT') else f'New Template {slot[:2]}-{variant}.docx')
            identity='builtin-'+slot
            with store.connect() as db:
                exists=db.execute('SELECT id FROM template_versions WHERE id=?',(identity,)).fetchone()
            if not exists:
                content=source.read_bytes();self._write(identity,content)
                with store.connect() as db:
                    db.execute('INSERT OR IGNORE INTO template_versions VALUES (?,?,?,?,?,?,?,?,?)',
                        (identity,slot,'Original DSENY template','legacy','{}',hashlib.sha256(content).hexdigest(),self.now(),1,0))
                    db.execute('INSERT OR IGNORE INTO template_active VALUES (?,?)',(slot,identity))

    @staticmethod
    def now(): return datetime.now(timezone.utc).isoformat()

    def _write(self,identity,content):
        target=self.directory/(identity+'.docx');tmp=target.with_suffix('.tmp');tmp.write_bytes(content);tmp.replace(target)

    def get(self,identity):
        with self.store.connect() as db:
            row=db.execute('SELECT * FROM template_versions WHERE id=?',(identity,)).fetchone()
        if not row: raise KeyError('Template version not found')
        result=dict(row);result['mapping']=json.loads(result['mapping'])
        path=self.directory/(result['id']+'.docx')
        if not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest()!=result['sha256']:
            raise ValueError('Template file is missing or its checksum changed. Restore the original upload.')
        result['path']=str(path)
        return result

    def public(self,record):
        content=Path(record['path']).read_bytes()
        info={k:v for k,v in record.items() if k!='path'}
        return dict(info,**inspect(content,record['mode'],record['mapping'],record['slot'],builtin=bool(record['builtin'])))

    def list(self):
        with self.store.connect() as db:
            ids=[r['id'] for r in db.execute('SELECT id FROM template_versions WHERE id NOT IN (SELECT id FROM template_archives) ORDER BY created_at DESC')]
            active={r['slot']:r['version_id'] for r in db.execute('SELECT * FROM template_active')}
        return {'active':active,'versions':[self.public(self.get(i)) for i in ids],
                'fields':list(TEMPLATE_FIELDS),'slots':list(SLOTS),
                'notice':'Template uploads use backend storage. On Render Free they can be lost on restart or deployment. Download every version; durable storage is not configured by this feature.'}

    def add(self,content,slot,label,mode,mapping):
        if slot not in SLOTS: raise ValueError('Unsupported report template slot')
        inspect(content,mode,mapping,slot)
        identity=uid();self._write(identity,content)
        with self.store.connect() as db:
            db.execute('INSERT INTO template_versions VALUES (?,?,?,?,?,?,?,?,?)',
                (identity,slot,label,mode,json.dumps(mapping),hashlib.sha256(content).hexdigest(),self.now(),0,0))
        return self.public(self.get(identity))

    def active(self):
        with self.store.connect() as db: pointers=list(db.execute('SELECT * FROM template_active'))
        return {r['slot']:self.get(r['version_id']) for r in pointers}

    def archive(self,identity):
        record=self.get(identity)
        with self.store.connect() as db:
            if record['builtin'] or db.execute('SELECT 1 FROM template_active WHERE version_id=?',(identity,)).fetchone():
                raise ValueError('Original and active templates cannot be archived; choose another active version first')
            db.execute('INSERT OR IGNORE INTO template_archives VALUES (?)',(identity,))
        return {'archived':True,'notice':'Hidden from the list; source and version metadata remain for report history.'}

    def activate(self,identity,reviewed):
        record=self.get(identity);info=self.public(record)
        if not reviewed: raise ValueError('Confirm that you reviewed the sample formatting and facts')
        if not info['valid']: raise ValueError('; '.join(info['errors']))
        if not record['previewed']: raise ValueError('Generate a successful sample for this exact template and mapping version first')
        with self.store.connect() as db:
            db.execute('DELETE FROM template_archives WHERE id=?',(identity,))
            db.execute('UPDATE template_active SET version_id=? WHERE slot=?',(identity,record['slot']))
        return {'active_version':identity,'slot':record['slot']}


def sample_report(slot):
    equipment=slot[:2];status=slot[3:]
    return {'equipment':equipment,'status':status,'client':'SAMPLE CLIENT','project':'SAMPLE PROJECT',
        'requested_by':'Sample requester','request_date':'2026-01-14','inspection_date':'2026-01-15',
        'permit':'SAMPLE PERMIT','address':'Sample site — not a real inspection','locations':'Sample north elevation, Level 2',
        'inspection_code':'I','drawings_on_site':'Yes','drawings_approved':'Yes',
        'attendees':[{'name':'SAMPLE INSPECTOR (not actual attendance)','company':'Sample company','title':'PE'}],
        'open_items':[{'location':'Sample platform edge','action':'Sample corrective action — not an actual finding'}] if status=='OIs' else [],
        'site':dict(report_type='Unknown / custom',other_name='SAMPLE OTHER ITEM',details='Sample inspection scope',**{k:'10' for k in ('length_ft','height_ft','sections','width_ft','tie_horizontal','tie_vertical',
            'planked_levels','outrigger_levels','tie_type','dunnage_type')}),
        'source_references':['SAMPLE ONLY — review formatting, not a certified inspection'], 'photo_references':[]}
