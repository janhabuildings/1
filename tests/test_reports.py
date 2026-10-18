from copy import deepcopy
import json
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile
from docx import Document
import report_agent as agent

BASE = json.loads((agent.ROOT/'examples/site.json').read_text())

class Reports(unittest.TestCase):
    def test_four_variants_and_preserved_package(self):
        for equipment in ('SF','SH'):
            for status in ('OIs','No OIs'):
                with self.subTest(equipment=equipment,status=status), tempfile.TemporaryDirectory() as temp:
                    d=deepcopy(BASE); d.update(equipment=equipment,status=status)
                    d['site'].update(sections='2',width_ft='12')
                    if status=='No OIs': d['open_items']=[]
                    out=Path(temp)/'report.docx'; master=agent.fill_report(d,out)
                    doc=Document(out)
                    text='\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
                    for expected in ('DEMO CLIENT','DEMO-PROJECT','DEMO-PERMIT','Example Inspector',agent.product(d),'100'):
                        self.assertIn(expected,text)
                    for stale in ('X01113434','270 East 167','Naseer','Kamruz','Noor Earpy','Mariusz','deteriorated mudsills','vertical plywood'):
                        self.assertNotIn(stale,text)
                    if equipment=='SF': self.assertNotIn('Sidewalk Shed',text)
                    self.assertEqual(doc.tables[3].cell(0,0).text,Document(master).tables[3].cell(0,0).text)
                    if status=='OIs':
                        self.assertIn(d['open_items'][0]['action'],text)
                        self.assertEqual(len(doc.tables[4].rows), 3 if equipment=='SF' else 4)
                    with ZipFile(master) as a, ZipFile(out) as b:
                        self.assertEqual(a.namelist(),b.namelist())
                        for name in a.namelist():
                            if name!='word/document.xml': self.assertEqual(a.read(name),b.read(name),name)
                        original=Document(master)
                        self.assertEqual(original.tables[1].cell(2,3)._tc.tcPr.xml,doc.tables[1].cell(2,3)._tc.tcPr.xml)
                        self.assertEqual(original.sections[0]._sectPr.xml,doc.sections[0]._sectPr.xml)
    def test_reject_unsafe_or_incomplete_inputs(self):
        for change in ({'status':'No OIs'}, {'open_items':[]}, {'recipients_confirmed':False},
                       {'inspection_date':'2026-02-30'}, {'to':['bad\r\nBcc: x@example.com']},
                       {'drawings_on_site':None}, {'attendees':[]}):
            with self.subTest(change=change):
                d=deepcopy(BASE); d.update(change)
                with self.assertRaises(ValueError): agent.validate(d)
    def test_subject(self):
        self.assertEqual(agent.subject(BASE),'2026-05-18 DEMO-PROJECT Supported Scaffold-I-Inspection-OIs')


class OfficeIntegration(unittest.TestCase):
    @unittest.skipUnless((agent.ROOT/'.tools/bin/libreoffice').exists() or agent.shutil.which('libreoffice'), 'LibreOffice is not installed')
    def test_four_email_outputs(self):
        from email import policy
        from email.parser import BytesParser
        for equipment in ('SF','SH'):
            for status in ('OIs','No OIs'):
                with self.subTest(equipment=equipment,status=status), tempfile.TemporaryDirectory() as temp:
                    d=deepcopy(BASE); d.update(equipment=equipment,status=status)
                    d['site'].update(sections='2',width_ft='12')
                    if status=='No OIs': d['open_items']=[]
                    manifest_path,m=agent.build(d,Path(temp))
                    self.assertEqual(m['status'],'local_draft_generated')
                    self.assertFalse(m['outlook_draft_created'])
                    mail=BytesParser(policy=policy.default).parsebytes((manifest_path.parent/m['eml']).read_bytes())
                    self.assertEqual(str(mail['Subject']),agent.subject(d))
                    html=mail.get_body(preferencelist=('html',)).get_content()
                    from lxml.html import fromstring
                    rendered_text=' '.join(fromstring(html).text_content().split())
                    attachments=list(mail.iter_attachments())
                    if status=='OIs':
                        self.assertIn('DEMO-PROJECT',html)
                        self.assertIn('Section 5',rendered_text)
                        self.assertIn('Replace damaged plank',rendered_text)
                        self.assertEqual(attachments,[])
                    else:
                        self.assertEqual(len(attachments),1)
                        self.assertEqual(attachments[0].get_content_type(),'application/pdf')
                        self.assertTrue(attachments[0].get_payload(decode=True).startswith(b'%PDF-'))
                        self.assertNotIn('signed and sealed',html)

if __name__=='__main__': unittest.main()
