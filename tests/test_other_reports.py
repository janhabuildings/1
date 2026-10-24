from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile
from docx import Document
from fastapi.testclient import TestClient
from backend.app import create_app
from backend.core import Settings
from backend.templates import sample_report,TemplateManager
import report_agent

ROOT=Path(__file__).resolve().parents[1]
HEADERS={'Authorization':'Bearer other-report-test-token'}


class OtherReports(unittest.TestCase):
    def test_full_template_general_sections_custom_name_and_preserved_styles(self):
        with tempfile.TemporaryDirectory() as directory:
            for status in ('OIs','No OIs'):
                data=sample_report('OT-'+status)
                data['site']={'report_type':'Unknown / custom','other_name':'Temporary Debris Screen','details':'North elevation protection'}
                path=Path(directory)/(status.replace(' ','')+'.docx')
                source=report_agent.fill_report(data,path,draft=True)
                doc=Document(path)
                text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
                self.assertIn('Temporary Debris Screen',text)
                self.assertIn('North elevation protection',text)
                self.assertIn('Section 3',doc.tables[3].cell(0,0).text)
                for stale in ('Scaffold Specific','Shed Specific','CUSTOM VERBIAGE','IF OPEN ITEM REPORT','IF NO OPEN ITEM REPORT',
                    'Requested DATE BY COMPANY','STEPS:','Kamruz','Dominic','catchall','scaffold or sheds','able to be occupied','Hoist Inspected'):
                    self.assertNotIn(stale,text)
                if status=='OIs':
                    self.assertIn('Sample corrective action',text)
                    self.assertIn('Section 5',text)
                else:
                    self.assertIn('No open items were reported',text)
                    self.assertNotIn('Section 5',text)
                with ZipFile(source) as a,ZipFile(path) as b:
                    for name in a.namelist():
                        if name!='word/document.xml':self.assertEqual(a.read(name),b.read(name),name)
                self.assertIn('Temporary Debris Screen',report_agent.subject(data))
            for kind in ('SOE','Fall Protection','Fence'):
                data['site']={'report_type':kind}
                self.assertEqual(report_agent.product(data),kind)
            data['site']={'report_type':'Unknown / custom'}
            with self.assertRaisesRegex(ValueError,'Other item name'):report_agent.validate(data,draft=True)

    def test_other_api_outputs_both_formats_and_template_slots(self):
        with tempfile.TemporaryDirectory() as directory:
            settings=Settings(Path(directory),'other-report-test-token',office_bin=str(ROOT/'.tools/bin/libreoffice'))
            app=create_app(settings)
            with TestClient(app) as client:
                catalog=client.get('/templates',headers=HEADERS).json()
                self.assertIn('OT-OIs',catalog['slots']);self.assertIn('OT-No OIs',catalog['slots'])
                for status in ('OIs','No OIs'):
                    report=sample_report('OT-'+status)
                    report['site']={'report_type':'Unknown / custom','other_name':'Temporary Debris Screen'}
                    response=client.post('/inspections',headers=HEADERS,json={'title':'Other inspection',
                        'project':report['project'],'inspection_date':report['inspection_date'],'reports':[report]})
                    self.assertEqual(response.status_code,201,response.text)
                    identity=response.json()['id']
                    result=client.post('/inspections/'+identity+'/reports',headers=HEADERS,json={})
                    self.assertEqual(result.status_code,200,result.text)
                    body=result.json();self.assertEqual(body['reports'][0]['status'],'local_draft_generated',body)
                    doc=next(a for a in body['artifacts'] if a['mime'].endswith('wordprocessingml.document'))
                    self.assertEqual(doc['equipment'],'OT')
                    content=client.get('/artifacts/'+doc['id'],headers=HEADERS).content
                    self.assertIn('Temporary Debris Screen',report_agent.physical_cells(Document(BytesIO(content)).tables[1].rows[3])[1].text)
                    if status=='No OIs':
                        pdf=next(a for a in body['artifacts'] if a['mime']=='application/pdf')
                        self.assertTrue(client.get('/artifacts/'+pdf['id'],headers=HEADERS).content.startswith(b'%PDF-'))
                        html=next(a for a in body['artifacts'] if a['mime']=='text/html')
                        self.assertNotIn('Certification for use',client.get('/artifacts/'+html['id'],headers=HEADERS).text)
                preview=client.post('/templates/builtin-OT-No OIs/preview',headers=HEADERS)
                self.assertTrue(preview.json()['preview_success'],preview.text)
