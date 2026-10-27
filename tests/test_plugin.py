import base64
from copy import deepcopy
from email import policy
from email.parser import BytesParser
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

ROOT=Path(__file__).resolve().parents[1]
PLUGIN=ROOT/'plugins/dseny-inspections'
spec=importlib.util.spec_from_file_location('plugin_runner',PLUGIN/'scripts/run_report.py')
runner=importlib.util.module_from_spec(spec);spec.loader.exec_module(runner)
EXAMPLE=json.loads((PLUGIN/'references/input-example.json').read_text())
OFFICE=ROOT/'.tools/bin/libreoffice'

class PluginInputs(unittest.TestCase):
    def test_pending_facts_do_not_become_na_or_a_named_inspector(self):
        data=deepcopy(EXAMPLE['reports'][0])
        data.pop('attendees');data.pop('drawings_on_site');data.pop('site')
        prepared,missing=runner.prepare(data)
        self.assertIn('inspecting_pe',missing)
        self.assertEqual(prepared['attendees'][0]['name'],runner.PENDING)
        self.assertEqual(prepared['drawings_on_site'],runner.PENDING)
        self.assertIn('site.tie_type',missing)
        self.assertNotEqual(prepared['site']['tie_type'],'NA')

    def test_missing_outcome_and_dates_are_not_inferred(self):
        for key in ['status','inspection_date','request_date','open_items']:
            d=deepcopy(EXAMPLE['reports'][0]);d.pop(key)
            with self.subTest(key=key),self.assertRaises(ValueError): runner.prepare(d)

    def test_no_ois_conflicts_with_finding(self):
        d=deepcopy(EXAMPLE['reports'][0]);d['status']='No OIs'
        with self.assertRaises(ValueError): runner.prepare(d)

    def test_duplicate_equipment_fails_before_partial_output(self):
        data=deepcopy(EXAMPLE)
        data['reports'].append(deepcopy(data['reports'][0]))
        data['reports'][-1]['inspection_code']='IX'
        with tempfile.TemporaryDirectory() as tmp:
            target=Path(tmp)/'out'
            with self.assertRaises(ValueError): runner.generate(data,target)
            self.assertFalse(target.exists())

    def test_subject_header_injection_is_rejected_before_writing(self):
        data=deepcopy(EXAMPLE)
        data['reports'][0]['subject_project']='Project\nBcc: bad@example.com'
        with tempfile.TemporaryDirectory() as tmp:
            target=Path(tmp)/'out'
            with self.assertRaises(ValueError): runner.generate(data,target)
            self.assertFalse(target.exists())

    def test_recipients_and_overwrite_protection(self):
        for mailbox in [{'account':'bad\nBcc:x@example.com'},
                        {'account':'janhabuildings@gmail.com','to':['client@example.com']}]:
            d=deepcopy(EXAMPLE);d['mailbox']=mailbox
            with tempfile.TemporaryDirectory() as tmp,self.assertRaises(ValueError):
                runner.generate(d,Path(tmp)/'out')
        with tempfile.TemporaryDirectory() as tmp,self.assertRaises(ValueError):
            runner.generate(EXAMPLE,Path(tmp))

    def test_converter_failure_retains_docx_without_false_draft(self):
        with tempfile.TemporaryDirectory() as tmp,patch.object(runner,'office_convert',side_effect=RuntimeError('converter unavailable')):
            manifest=runner.generate(EXAMPLE,Path(tmp)/'out')
            self.assertFalse(manifest['gmail_drafts_created'])
            for report in manifest['reports']:
                self.assertEqual(report['status'],'conversion_blocked')
                self.assertTrue(Path(report['docx']).is_file())
                self.assertNotIn('gmail_handoff',report)

    def test_word_branch_is_explicitly_platform_limited(self):
        if sys.platform=='win32': self.skipTest('Windows branch requires real Word')
        with tempfile.TemporaryDirectory() as tmp:
            m=runner.generate(EXAMPLE,Path(tmp)/'word',renderer='word')
            self.assertTrue(all(r['status']=='conversion_blocked' for r in m['reports']))
            self.assertTrue(all('Windows' in r['error'] for r in m['reports']))

class PluginPackage(unittest.TestCase):
    def test_extracted_bundle_is_self_contained(self):
        completed=subprocess.run([sys.executable,str(ROOT/'scripts/build_plugin.py')],check=True,capture_output=True,text=True)
        package=Path(completed.stdout.strip())
        with tempfile.TemporaryDirectory() as tmp:
            with ZipFile(package) as z:
                names=z.namelist()
                self.assertFalse(any('__pycache__' in n or '/output/' in n for n in names))
                self.assertIn('.agents/plugins/marketplace.json',names)
                z.extractall(tmp)
            root=Path(tmp);plugin=root/'plugins/dseny-inspections'
            manifest=json.loads((plugin/'.codex-plugin/plugin.json').read_text())
            for key in ('skills','apps'):
                self.assertTrue((plugin/manifest[key]).exists())
            app=json.loads((plugin/'.app.json').read_text())
            self.assertEqual(app['apps']['gmail']['id'],'connector_2128aebfecb84f64a069897515042a44')
            checks=json.loads((plugin/'checksums.json').read_text())
            for name,digest in checks.items(): self.assertEqual(hashlib.sha256((plugin/name).read_bytes()).hexdigest(),digest)
            self.assertEqual(len(list((plugin/'runtime/templates').glob('*.docx'))),7)
            if not OFFICE.exists(): self.skipTest('LibreOffice not available for extracted runtime integration')
            env=dict(os.environ,DSENY_OFFICE_BIN=str(OFFICE))
            output=root/'state/out'
            run=subprocess.run([sys.executable,str(plugin/'scripts/run_report.py'),str(plugin/'references/input-example.json'),
                                '--output',str(output)],env=env,capture_output=True,text=True,timeout=120)
            self.assertEqual(run.returncode,0,run.stderr)
            m=json.loads((output/'manifest.json').read_text())
            self.assertEqual(len(m['reports']),2)
            self.assertFalse(m['gmail_drafts_created'])
            for report in m['reports']:
                self.assertEqual(report['status'],'local_draft_generated')
                self.assertFalse(report['delivery_ready'])
                handoff=json.loads(Path(report['gmail_handoff']).read_text())
                self.assertEqual(handoff['account'],'janhabuildings@gmail.com')
                self.assertEqual(handoff['to'],[])
                self.assertIn('(ALTERNATE)',handoff['subject'])
                raw=json.loads(Path(handoff['raw_mime_json_path']).read_text())
                message=BytesParser(policy=policy.default).parsebytes(base64.urlsafe_b64decode(raw['message']['raw']))
                self.assertEqual(str(message['Subject']),handoff['subject'])
                self.assertIsNone(message['To'])
                self.assertEqual(str(message['From']),'janhabuildings@gmail.com')
                self.assertIn('DEMO-PROJECT',handoff['body_html'] if report['equipment']=='SF' else handoff['subject'])
                if report['equipment']=='SF':
                    from lxml.html import fromstring
                    text=' '.join(fromstring(handoff['body_html']).text_content().split())
                    self.assertIn('4–5',text)
                    self.assertIn('Install missing toeboards',text)
                    self.assertIn('Section 5',text)
                    self.assertFalse(handoff['attachments'])
                    self.assertEqual(list(message.iter_attachments()),[])
                else:
                    self.assertEqual(len(handoff['attachments']),1)
                    attachment=handoff['attachments'][0]
                    contents=Path(attachment['path']).read_bytes()
                    self.assertTrue(contents.startswith(b'%PDF-'))
                    self.assertEqual(hashlib.sha256(contents).hexdigest(),attachment['sha256'])
                    self.assertEqual(list(message.iter_attachments())[0].get_payload(decode=True),contents)
                    self.assertNotIn('signed and sealed',handoff['body_html'])

class EmailSpacing(unittest.TestCase):
    def test_docx_spacing_is_inlined_without_stretching_tables_or_changing_words(self):
        from docx import Document
        from docx.shared import Pt
        from lxml import html
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);doc=Document()
            first=doc.add_paragraph('Keep  these words together.')
            first.paragraph_format.space_before=Pt(6);first.paragraph_format.space_after=Pt(9);first.paragraph_format.line_spacing=1.5
            cell=doc.add_table(rows=1,cols=1).cell(0,0)
            cell.text='Table cell';cell.paragraphs[0].paragraph_format.space_before=Pt(0);cell.paragraphs[0].paragraph_format.space_after=Pt(0);cell.paragraphs[0].paragraph_format.line_spacing=1
            doc.save(root/'source.docx')
            source='<html><head><style>p {margin-bottom:12pt;line-height:115%}</style></head><body><p>Keep &nbsp;these <b>words</b> together.</p><table width="648"><tr><td style="padding:0in 0.08in"><p>Table cell</p></td></tr></table></body></html>'
            (root/'source.html').write_text(source)
            result=html.fromstring(runner.email_html(root/'source.html',root/'source.docx'))
            first=result.xpath('//p')[0];css=first.get('style')
            self.assertIn('margin-top:6pt;',css);self.assertIn('margin-bottom:9pt;',css);self.assertIn('line-height:1.5;',css)
            self.assertIn('Keep \xa0these words together.',first.text_content())
            self.assertEqual(first.find('b').text,'words')
            self.assertIn('margin-top:0pt;',result.xpath('//td/p')[0].get('style'))
            self.assertIn('margin-bottom:0pt;',result.xpath('//td/p')[0].get('style'))
            self.assertIn('line-height:1;',result.xpath('//td/p')[0].get('style'))
            table=result.xpath('//table')[0];self.assertEqual(table.get('width'),'648');self.assertEqual(table.get('cellspacing'),'0')
            self.assertIn('width:648px;',table.get('style'));self.assertNotRegex(table.get('style'),r'(?:^|;)width:100%;')
            self.assertEqual(result.xpath('//td')[0].get('style'),'padding:0in 0.08in')
            self.assertIn('width:648px;',result.xpath('//div[@class="report-content"]')[0].get('style'))

    def test_converter_spacing_is_inline_when_no_docx_is_supplied(self):
        from lxml import html
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'sample.html'
            path.write_text('<html><head><style>p {line-height:115%;margin-bottom:8pt}</style></head><body><p style="margin-bottom:0in;line-height:100%">Sentence</p><p>Table spacing</p></body></html>')
            result=html.fromstring(runner.email_html(path))
            self.assertIn('margin-top:0;',result.xpath('//p')[0].get('style'))
            self.assertIn('line-height:100%;',result.xpath('//p')[0].get('style'))
            self.assertIn('margin-bottom:8pt;',result.xpath('//p')[1].get('style'))

if __name__=='__main__': unittest.main()
