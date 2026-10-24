from io import BytesIO
from pathlib import Path
import tempfile
import unittest
import uuid
from email import policy
from email.parser import BytesParser
from PIL import Image,ImageDraw
from pypdf import PdfReader
from docx import Document
from fastapi.testclient import TestClient
from backend.app import create_app
from backend.core import Settings
from backend.templates import sample_report
import report_agent

ROOT=Path(__file__).resolve().parents[1]
HEADERS={'Authorization':'Bearer updates-test-token'}


def test_image():
    im=Image.new('RGB',(250,100),'white');ImageDraw.Draw(im).text((12,30),'TEST IMAGE ONLY',fill='black')
    out=BytesIO();im.save(out,format='PNG');return out.getvalue()


class CellMapping(unittest.TestCase):
    def test_labels_and_following_values_in_all_six_variants(self):
        with tempfile.TemporaryDirectory() as directory:
            for equipment in ('SF','SH','OT'):
                for status in ('OIs','No OIs'):
                    data=sample_report(equipment+'-'+status);path=Path(directory)/(equipment+status+'.docx')
                    report_agent.fill_report(data,path,draft=True);doc=Document(path);table=doc.tables[1]
                    cells=report_agent.physical_cells
                    for row,label in ((1,'Client:'),(2,'Request date:'),(3,'Product Inspected:'),(4,'Permit#:'),(5,'Address:'),(6,'PE Drawings on site?'),(7,'Drawings DOB/SCA approved?')):
                        self.assertEqual(cells(table.rows[row])[0].text.strip(),label,(equipment,status,row))
                    self.assertEqual(cells(table.rows[1])[1].text,data['client'])
                    self.assertEqual(cells(table.rows[1])[2].text.strip(),'Project:')
                    self.assertEqual(cells(table.rows[1])[3].text,data['project'])
                    self.assertEqual(cells(table.rows[2])[1].text,'01.14.2026')
                    self.assertEqual(cells(table.rows[2])[3].text,'01.15.2026')
                    self.assertEqual(cells(table.rows[2])[5].text,data['requested_by'])
                    self.assertEqual(cells(table.rows[3])[1].text,report_agent.product(data))
                    self.assertEqual(cells(table.rows[6])[1].text,'Yes')
                    for section in doc.tables[2:5]:
                        previous=section._element.getprevious();self.assertEqual(previous.tag,report_agent.W+'p')
                        self.assertEqual(previous.getprevious().tag,report_agent.W+'p')
                    if status=='OIs':
                        number=doc.tables[4].rows[-1].cells[0]
                        self.assertEqual(number.text,'1')
                        self.assertEqual(number._tc.xpath('.//w:numId')[0].get(report_agent.W+'val'),'0')
                    if equipment=='SF':
                        site=doc.tables[3];location=next(i for i,r in enumerate(site.rows) if r.cells[0].text.strip()=='LOCATION/S:')
                        self.assertEqual(cells(site.rows[location+2])[0].text,'Planked Levels')
                        self.assertEqual(cells(site.rows[location+2])[1].text,'10')
                        self.assertEqual(cells(site.rows[location+2])[2].text,'Planked Outrigger Levels')
                        self.assertEqual(cells(site.rows[location+2])[3].text,'10')

    def test_scaffold_hoist_chute_fields_are_retained_in_both_variants(self):
        for status in ('OIs','No OIs'):
            data=sample_report('SF-'+status)
            data['site'].update(hoist_inspected='Yes',chute_inspected='Yes',hoist_locations='North courtyard',hoist_height='48',hoist_model='Example Hoist',chute_locations='East bay',chute_height='36',chute_model='Example Chute')
            with tempfile.TemporaryDirectory() as directory:
                path=Path(directory)/'report.docx';report_agent.fill_report(data,path,draft=True)
                table=Document(path).tables[3]
                text=' '.join(c.text for row in table.rows for c in report_agent.physical_cells(row))
                for expected in ('Hoist Inspected?','Chute Inspected?','North courtyard','48','Example Hoist','East bay','36','Example Chute'):
                    self.assertIn(expected,text)


class LifecycleAndStamp(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.app=create_app(Settings(Path(self.temp.name),'updates-test-token',office_bin=str(ROOT/'.tools/bin/libreoffice')))
        self.client=TestClient(self.app)
    def tearDown(self):self.client.close();self.temp.cleanup()
    def inspection(self,status='No OIs'):
        report=sample_report('OT-'+status);report['attendees'][0]['title']='P.E.'
        result=self.client.post('/inspections',headers=HEADERS,json={'title':'Sync fixture','project':report['project'],
            'inspection_date':report['inspection_date'],'reports':[report]})
        self.assertEqual(result.status_code,201,result.text);return result.json()
    def test_deletions_artifact_sharing_and_no_resurrection(self):
        item=self.inspection();identity=item['id'];asset=str(uuid.uuid4())
        upload_headers=dict(HEADERS,**{'Content-Type':'image/png','X-Asset-ID':asset})
        result=self.client.post('/inspections/'+identity+'/media?kind=photo&filename=test.png',headers=upload_headers,content=test_image())
        self.assertEqual(result.status_code,201,result.text)
        self.assertEqual(len(self.client.get('/inspections/'+identity+'/media',headers=HEADERS).json()),1)
        self.assertEqual(self.client.delete('/media/'+asset).status_code,401)
        self.assertEqual(self.client.delete('/media/'+asset,headers=HEADERS).status_code,200)
        self.assertIn(asset,self.client.get('/sync/deletions',headers=HEADERS).json()['media'])
        self.assertEqual(self.client.post('/inspections/'+identity+'/media?kind=photo&filename=test.png',headers=upload_headers,content=test_image()).status_code,409)
        generated=self.client.post('/inspections/'+identity+'/reports',headers=HEADERS,json={})
        self.assertEqual(generated.status_code,200,generated.text)
        shared=self.client.get('/inspections/'+identity+'/artifacts',headers=HEADERS).json()
        self.assertTrue(shared['artifacts'])
        file=shared['artifacts'][0]['id'];self.assertEqual(self.client.delete('/artifacts/'+file,headers=HEADERS).status_code,200)
        self.assertEqual(self.client.get('/artifacts/'+file,headers=HEADERS).status_code,404)
        self.assertEqual(self.client.delete('/inspections/'+identity,headers=HEADERS).status_code,200)
        self.assertEqual(self.client.get('/inspections',headers=HEADERS).json(),[])
        self.assertEqual(self.client.put('/inspections/'+identity,headers=HEADERS,json=item).status_code,409)
        self.assertIn(identity,self.client.get('/sync/deletions',headers=HEADERS).json()['inspections'])
    def test_compact_no_ois_pagination_with_complete_scaffold_and_reused_stamp(self):
        owner='Sample PE'
        stamp=self.client.post('/signing-assets',params={'owner':owner},headers=HEADERS,content=test_image()).json()['id']
        for equipment in ('SH','SF'):
            report=sample_report(equipment+'-No OIs')
            report['attendees']=[{'name':'Example Contact','company':'Contractor','title':'Competent Person'},{'name':'Example Inspector','company':'NYC SCA','title':'Safety Inspector'},{'name':owner,'company':'DSENY','title':'PE'}]
            report['site'].update(hoist_inspected='No',chute_inspected='No')
            result=self.client.post('/inspections',headers=HEADERS,json={'title':'Layout sample','project':report['project'],'inspection_date':report['inspection_date'],'reports':[report]}).json()
            generated=self.client.post('/inspections/'+result['id']+'/reports',headers=HEADERS,json={'stamp_id':stamp,'stamp_reviewed':True,'mailbox_email':'inspector@example.com'}).json()
            pdf=next(a for a in generated['artifacts'] if a['mime']=='application/pdf')
            reader=PdfReader(BytesIO(self.client.get('/artifacts/'+pdf['id'],headers=HEADERS).content))
            self.assertEqual(len(reader.pages),2)
            self.assertIn('Section 2',' '.join(reader.pages[0].extract_text().split()))
            self.assertNotIn('Section 3',reader.pages[0].extract_text())
            self.assertIn('Section 3',' '.join(reader.pages[1].extract_text().split()))
            self.assertIn('DSENY adheres',reader.pages[1].extract_text())
            self.assertIn('PE signature/seal:',reader.pages[1].extract_text())
            if equipment=='SF':self.assertIn('Hoist Model',reader.pages[1].extract_text())
            eml=next(a for a in generated['artifacts'] if a['filename'].endswith('.eml'))
            message=BytesParser(policy=policy.default).parsebytes(self.client.get('/artifacts/'+eml['id'],headers=HEADERS).content)
            self.assertEqual(str(message['From']),'inspector@example.com')

    def test_stamp_review_owner_page_two_and_email_attachment(self):
        item=self.inspection();identity=item['id'];owner=item['reports'][0]['attendees'][0]['name']
        self.assertEqual(self.client.post('/signing-assets?owner=Test',content=test_image()).status_code,401)
        stamp=self.client.post('/signing-assets',params={'owner':owner},headers=HEADERS,content=test_image())
        self.assertEqual(stamp.status_code,200,stamp.text);asset=stamp.json()['id']
        self.assertEqual(self.client.post('/inspections/'+identity+'/reports',headers=HEADERS,json={'stamp_id':asset}).status_code,422)
        result=self.client.post('/inspections/'+identity+'/reports',headers=HEADERS,json={'stamp_id':asset,'stamp_reviewed':True})
        self.assertEqual(result.status_code,200,result.text);body=result.json()
        self.assertEqual(body['reports'][0]['visual_stamp']['page'],2)
        self.assertFalse(body['reports'][0]['visual_stamp']['cryptographic_signature'])
        pdf=next(a for a in body['artifacts'] if a['mime']=='application/pdf')
        pdf_bytes=self.client.get('/artifacts/'+pdf['id'],headers=HEADERS).content
        reader=PdfReader(BytesIO(pdf_bytes));self.assertGreaterEqual(len(reader.pages),2)
        self.assertIn('PE signature/seal:',reader.pages[1].extract_text())
        self.assertNotIn('PE signature/seal:',reader.pages[0].extract_text())
        self.assertTrue(len(list(reader.pages[1].images)))
        eml=next(a for a in body['artifacts'] if a['filename'].endswith('.eml'))
        message=BytesParser(policy=policy.default).parsebytes(self.client.get('/artifacts/'+eml['id'],headers=HEADERS).content)
        self.assertEqual(next(message.iter_attachments()).get_payload(decode=True),pdf_bytes)
        docx=next(a for a in body['artifacts'] if a['filename'].endswith('.docx'))
        doc=Document(BytesIO(self.client.get('/artifacts/'+docx['id'],headers=HEADERS).content))
        self.assertEqual(len(doc.tables[2].rows),2,'P.E. must not create another pending PE')
        wrong=self.client.post('/signing-assets',params={'owner':'Another engineer'},headers=HEADERS,content=test_image()).json()['id']
        self.assertEqual(self.client.post('/inspections/'+identity+'/reports',headers=HEADERS,json={'stamp_id':wrong,'stamp_reviewed':True}).status_code,422)
        self.assertEqual(self.client.delete('/signing-assets/'+asset,headers=HEADERS).status_code,200)
