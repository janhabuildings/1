from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile
import hashlib
import unittest
from docx import Document
from backend.templates import sample_report
from report_agent import fill_report,physical_cells,NS,subject

ROOT=Path(__file__).resolve().parents[1]

class LatestMaster(unittest.TestCase):
    def test_all_variants_use_master_styles_and_omit_preparation_text(self):
        master=ROOT/'templates/New Template Full.docx'
        with ZipFile(master) as source:
            style_parts={name:source.read(name) for name in ('word/styles.xml','word/theme/theme1.xml','word/fontTable.xml')}
        for equipment in ('SF','SH','OT'):
            for status in ('OIs','No OIs'):
                with self.subTest(equipment=equipment,status=status),TemporaryDirectory() as directory:
                    data=sample_report(equipment+'-'+status);path=Path(directory)/'report.docx'
                    selected=fill_report(data,path,draft=True);doc=Document(path)
                    self.assertEqual(doc.core_properties.identifier,'DSENY-2026-10-25')
                    with ZipFile(path) as output:
                        for part,content in style_parts.items():self.assertEqual(output.read(part),content)
                    self.assertEqual(physical_cells(doc.tables[0].rows[0])[1].text,'30-01 39 Avenue Long Island City NY 11101\n347.730.6990\xa0\xa0\xa0\xa0\xa0\xa0\xa0\xa0\xa0 inspections@dseny.nyc')
                    text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in physical_cells(row))
                    for unwanted in ('Vinny Madia','Tigran Avetisyan','Keltech','G Washington','START OF','END OF','IF CCI REPORT','CUSTOM VERBIAGE','DELETE ALL TEXT','STEPS:'):
                        self.assertNotIn(unwanted,text)
                    self.assertEqual(physical_cells(doc.tables[1].rows[0])[0].paragraphs[0].runs[0].font.name,'Aptos')
                    self.assertEqual(str(physical_cells(doc.tables[1].rows[0])[0].paragraphs[0].runs[0].font.color.rgb),'000000')
                    if status=='No OIs':
                        self.assertIn('CERTIFICATE OF COMPLETED INSTALLATION (CCI)',text)
                        self.assertTrue(subject(data).endswith('Inspection-PE CERT OF USE'))
                        self.assertNotIn('Section 5',text)
                        # New 12pt default is retained rather than legacy forced 10pt body.
                        defaults=doc.styles.element.find('w:docDefaults/w:rPrDefault/w:rPr/w:sz',NS)
                        self.assertEqual(defaults.get('{'+NS['w']+'}val'),'24')
                    else:self.assertIn('Section 5',text)

    def test_legacy_scaffold_variant_still_receives_correct_hoist_chute_rows(self):
        data=sample_report('SF-No OIs')
        data['site'].update(hoist_inspected='Yes',chute_inspected='Yes',hoist_locations='North',hoist_height='48',hoist_model='Example hoist',chute_locations='East',chute_height='36',chute_model='Example chute')
        with TemporaryDirectory() as directory:
            root=Path(directory);legacy=Document(ROOT/'templates/New Template SF-N-Ois.docx')
            legacy.core_properties.identifier=''
            for row in list(legacy.tables[3].rows)[6:]:legacy.tables[3]._tbl.remove(row._tr)
            legacy.tables[4]._element.getparent().remove(legacy.tables[4]._element)
            path=root/'legacy.docx';legacy.save(path)
            output=root/'report.docx';fill_report(data,output,draft=True,template=path)
            text=' '.join(c.text for row in Document(output).tables[3].rows for c in physical_cells(row))
            for term in ('Hoist Inspected?','Hoist Model','Example hoist','Chute Model','Example chute'):self.assertIn(term,text)
