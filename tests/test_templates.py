from io import BytesIO
import hashlib
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch,AsyncMock
from zipfile import ZipFile,ZIP_DEFLATED
from docx import Document
from docx.shared import Pt,RGBColor
from fastapi.testclient import TestClient
from backend.app import create_app
from backend.core import Settings,Store
from backend.templates import TemplateManager,sample_report
import report_agent

ROOT=Path(__file__).resolve().parents[1]
HEADERS={'Authorization':'Bearer template-test-token'}


def placeholder_doc(alias=False):
    doc=Document()
    p=doc.add_paragraph();a=p.add_run('{{pro');a.bold=True;a.font.color.rgb=RGBColor.from_string('008080');a.font.size=Pt(18)
    p.add_run('ject}}')
    for key in ('client','product','inspection_date','status','attendees','locations','open_items'):
        doc.add_paragraph('{{'+key+'}}')
    if alias: doc.add_paragraph('{{school_name}}')
    doc.sections[0].footer.paragraphs[0].text='Footer {{project}}'
    stream=BytesIO();doc.save(stream);return stream.getvalue()


class TemplateEngine(unittest.TestCase):
    def test_split_runs_headers_multiline_and_untouched_parts(self):
        with tempfile.TemporaryDirectory() as directory:
            p=Path(directory);source=p/'template.docx';source.write_bytes(placeholder_doc())
            data=sample_report('SF-OIs');data['open_items'].append({'location':'Another bay','action':'Second action'})
            report_agent.fill_report(data,p/'result.docx',draft=True,template=source,template_mode='placeholders')
            doc=Document(p/'result.docx')
            self.assertEqual(doc.paragraphs[0].text,'SAMPLE PROJECT')
            self.assertTrue(doc.paragraphs[0].runs[0].bold)
            self.assertEqual(str(doc.paragraphs[0].runs[0].font.color.rgb),'008080')
            text='\n'.join(p.text for p in doc.paragraphs)
            self.assertIn('2. Another bay: Second action',text)
            self.assertNotIn('{{',text)
            self.assertIn('Footer SAMPLE PROJECT',doc.sections[0].footer.paragraphs[0].text)
            with ZipFile(source) as a,ZipFile(p/'result.docx') as b:
                for name in a.namelist():
                    if name not in ('word/document.xml','word/footer1.xml'):self.assertEqual(a.read(name),b.read(name),name)

    def test_malformed_markers_and_external_content_fields(self):
        from docx.oxml import OxmlElement
        with tempfile.TemporaryDirectory() as directory:
            manager=TemplateManager(Store(directory))
            document=Document(BytesIO(placeholder_doc()));document.add_paragraph('{{invalid marker}}')
            buffer=BytesIO();document.save(buffer)
            record=manager.add(buffer.getvalue(),'SF-OIs','Malformed','placeholders',{})
            self.assertFalse(record['valid'])
            self.assertTrue(any('Malformed' in e for e in record['errors']))
            document=Document(BytesIO(placeholder_doc()))
            field=OxmlElement('w:instrText');field.text='DDEAUTO fixture'
            document.add_paragraph()._p.append(field)
            buffer=BytesIO();document.save(buffer)
            with self.assertRaisesRegex(ValueError,'External-content'):
                manager.add(buffer.getvalue(),'SF-OIs','External','placeholders',{})

    def test_reject_active_content_and_immutable_mapping_revisions(self):
        with tempfile.TemporaryDirectory() as directory:
            manager=TemplateManager(Store(directory));content=placeholder_doc(alias=True)
            record=manager.add(content,'SF-OIs','First','placeholders',{})
            self.assertFalse(record['valid'])
            revised=manager.add(content,'SF-OIs','Mapped','placeholders',{'school_name':'project'})
            self.assertTrue(revised['valid'])
            self.assertNotEqual(record['id'],revised['id'])
            self.assertEqual(manager.get(record['id'])['mapping'],{})
            with self.assertRaisesRegex(ValueError,'sample'): manager.activate(revised['id'],True)
            malicious=BytesIO()
            with ZipFile(BytesIO(content)) as original,ZipFile(malicious,'w',ZIP_DEFLATED) as z:
                for entry in original.infolist():z.writestr(entry,original.read(entry.filename))
                z.writestr('word/vbaProject.bin',b'fixture')
            with self.assertRaisesRegex(ValueError,'Macros'): manager.add(malicious.getvalue(),'SF-OIs','Bad','placeholders',{})
            with self.assertRaises(ValueError): manager.add(b'not a docx','SF-OIs','Bad','legacy',{})
            # A changed file cannot inherit a successful preview or activation.
            with manager.store.connect() as db: db.execute('UPDATE template_versions SET previewed=1 WHERE id=?',(revised['id'],))
            Path(manager.get(revised['id'])['path']).write_bytes(b'tampered')
            with self.assertRaisesRegex(ValueError,'checksum'): manager.activate(revised['id'],True)


class TemplateAPI(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.app=create_app(Settings(Path(self.temp.name),'template-test-token',office_bin=str(ROOT/'.tools/bin/libreoffice')))
        self.client=TestClient(self.app)
    def tearDown(self):self.client.close();self.temp.cleanup()

    def upload(self,content=None,slot='SF-OIs',mode='placeholders'):
        return self.client.post('/templates',params={'slot':slot,'mode':mode,'label':'Test revision'},
            headers=HEADERS,content=content or placeholder_doc(alias=True))

    def test_auth_mapping_preview_activation_generation_and_rollback(self):
        self.assertEqual(self.client.get('/templates').status_code,401)
        self.assertEqual(self.client.post('/templates?slot=SF-OIs',content=placeholder_doc()).status_code,401)
        versions=self.client.get('/templates',headers=HEADERS).json()
        self.assertEqual(len(versions['versions']),6)
        self.assertEqual(self.client.get('/inspections',headers=HEADERS).json(),[])
        response=self.upload();self.assertEqual(response.status_code,201,response.text)
        old=response.json()['id']
        response=self.client.post(f'/templates/{old}/mapping',headers=HEADERS,json={'mapping':{'school_name':'project'}})
        self.assertEqual(response.status_code,201,response.text);version=response.json()['id']
        self.assertNotEqual(version,old)
        self.assertEqual(self.client.post(f'/templates/{version}/activate',headers=HEADERS,json={'reviewed':True}).status_code,422)
        preview=self.client.post(f'/templates/{version}/preview',headers=HEADERS)
        self.assertEqual(preview.status_code,200,preview.text)
        self.assertTrue(preview.json()['preview_success'],preview.text)
        self.assertEqual(self.client.get('/inspections',headers=HEADERS).json(),[])
        doc=next(a for a in preview.json()['artifacts'] if a['filename'].endswith('.docx'))
        result=Document(BytesIO(self.client.get('/artifacts/'+doc['id'],headers=HEADERS).content))
        self.assertIn('SAMPLE PROJECT','\n'.join(p.text for p in result.paragraphs))
        self.assertEqual(self.client.get('/artifacts/'+doc['id']).status_code,401)
        self.assertEqual(self.client.post(f'/templates/{version}/activate',headers=HEADERS,json={}).status_code,422)
        self.assertEqual(self.client.post(f'/templates/{version}/activate',headers=HEADERS,json={'reviewed':True}).status_code,200)
        # A fresh app instance retains activation and history in persistent data directories.
        restarted=TemplateManager(Store(self.temp.name))
        self.assertEqual(restarted.active()['SF-OIs']['id'],version)
        report=sample_report('SF-OIs');report['project']='ACTUAL FIXTURE PROJECT'
        inspection=self.client.post('/inspections',headers=HEADERS,json={'title':'Test','project':report['project'],
            'inspection_date':report['inspection_date'],'reports':[report]}).json()
        generated=self.client.post('/inspections/'+inspection['id']+'/reports',headers=HEADERS,json={})
        self.assertEqual(generated.status_code,200,generated.text)
        self.assertEqual(generated.json()['reports'][0]['template_version']['id'],version)
        artifact=next(a for a in generated.json()['artifacts'] if a['filename'].endswith('.docx'))
        actual=Document(BytesIO(self.client.get('/artifacts/'+artifact['id'],headers=HEADERS).content))
        self.assertEqual(actual.paragraphs[0].text,'ACTUAL FIXTURE PROJECT')
        # Restoring the original also requires reviewing its own successful sample.
        original='builtin-SF-OIs'
        self.assertTrue(self.client.post(f'/templates/{original}/preview',headers=HEADERS).json()['preview_success'])
        self.assertEqual(self.client.post(f'/templates/{original}/activate',headers=HEADERS,json={'reviewed':True}).status_code,200)
        self.assertEqual(restarted.active()['SF-OIs']['id'],original)
        self.assertEqual(self.client.get('/templates/'+version+'/metadata',headers=HEADERS).json()['mapping'],{'school_name':'project'})
        downloaded=self.client.get('/templates/'+version+'/download',headers=HEADERS).content
        self.assertEqual(hashlib.sha256(downloaded).hexdigest(),restarted.get(version)['sha256'])

    def test_failed_sample_conversion_cannot_be_activated(self):
        version=self.upload(placeholder_doc()).json()['id']
        failed={'reports':[{'status':'docx_generated','error':'Fixture converter failure'}]}
        with patch('backend.app.asyncio.to_thread',new=AsyncMock(return_value=failed)):
            result=self.client.post(f'/templates/{version}/preview',headers=HEADERS)
        self.assertEqual(result.status_code,200,result.text)
        self.assertFalse(result.json()['preview_success'])
        self.assertEqual(self.client.post(f'/templates/{version}/activate',headers=HEADERS,json={'reviewed':True}).status_code,422)

    def test_legacy_no_open_items_preview_and_safety_checks(self):
        wrong=(ROOT/'templates/New Template SF-Ois.docx').read_bytes()
        self.assertEqual(self.upload(wrong,slot='SH-OIs',mode='legacy').status_code,422)
        self.assertEqual(self.upload(wrong,slot='SF-No OIs',mode='legacy').status_code,422)
        original='builtin-SH-No OIs'
        result=self.client.post(f'/templates/{original}/preview',headers=HEADERS)
        self.assertEqual(result.status_code,200,result.text)
        self.assertTrue(result.json()['preview_success'],result.text)
        pdf=next(a for a in result.json()['artifacts'] if a['filename'].endswith('.pdf'))
        self.assertTrue(self.client.get('/artifacts/'+pdf['id'],headers=HEADERS).content.startswith(b'%PDF-'))
        self.assertEqual(self.client.post('/templates?slot=other',headers=HEADERS,content=placeholder_doc()).status_code,422)
        self.assertEqual(self.client.post('/templates?slot=SF-OIs&filename=bad.docm',headers=HEADERS,content=placeholder_doc()).status_code,422)
